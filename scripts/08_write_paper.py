# -*- coding: utf-8 -*-
"""
08_write_paper.py —— 按学校模板生成论文初稿（Word + Markdown）

输入模板：
    D:\\大数据分析及数据可视化\\大作业\\5、《商业数据分析》 - 期末大作业模板.docx
输出：
    D:\\大数据分析及数据可视化\\大作业\\Olist商业数据分析_论文初稿.docx
    docs/paper_draft.md

用法：
    python scripts/08_write_paper.py
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(os.environ.get(
    "PAPER_TEMPLATE",
    r"D:\大数据分析及数据可视化\大作业\5、《商业数据分析》 - 期末大作业模板.docx",
))
OUT_DOCX = Path(os.environ.get(
    "PAPER_OUT",
    r"D:\大数据分析及数据可视化\大作业\Olist商业数据分析_论文初稿.docx",
))
OUT_MD = ROOT / "docs" / "paper_draft.md"

FIG_T = ROOT / "outputs" / "figures" / "templates"
FIG_B = ROOT / "outputs" / "figures" / "business"
DASH = ROOT / "outputs" / "figures" / "dashboard_preview.png"

TITLE = "Olist 巴西电商客户分层与销售预测"
AUTHOR = "黄丹彤"
COLLEGE = "（嘉兴大学 数据科学学院）"
DATE = "2026 年 11 月"

ABSTRACT = (
    "本文以 Olist 巴西电商公开数据集为研究对象，基于 2016—2018 年的 99,441 个订单、"
    "112,650 条商品明细和 99,224 条评价，构建从数据获取、清洗治理、商业建模到可视化大屏的"
    "端到端分析链路。在完成多表关联、缺失值处理和指标派生后，使用 RFM 八分层、同期群留存、"
    "ABC/Pareto、波士顿矩阵、SARIMA 和订单漏斗模型，识别客户价值结构、复购规律、品类贡献与"
    "销售趋势；同时复现教材 15 张模板图、绘制 10 张业务图并搭建 Streamlit 交互式大屏。"
    "研究发现：约 47.5% 的高价值客户贡献 78.0% 的 GMV；首购后第 1 个月平均留存率仅 0.44%；"
    "25.9% 的商品贡献 80.0% 的 GMV；延迟 4 天以上订单的平均评分由 4.29 降至 2.32；"
    "ARIMA(1,1,1) 的六个月回测 MAPE 为 5.59%。研究结果可为电商客户运营、品类管理、"
    "物流优化与销售预测提供参考。"
)
KEYWORDS = "电商数据分析；RFM 模型；同期群分析；波士顿矩阵；SARIMA；可视化大屏"


def B(text: str):
    return ("p", text)


def T(caption: str, headers: list[str], rows: list[list[str]],
      widths: list[float] | None = None, font_size: float = 10.5):
    return ("table", caption, headers, rows, widths, font_size)


def F(path: Path, caption: str, width: float = 14.5):
    return ("figure", path, caption, width)


HEADING_CLEAN = {
    " 引言": "引言",
    "数据源介绍（": "数据源介绍",
    "合规获取方式（": "合规获取方式",
    "数据来源合规声明（": "数据来源合规声明",
    "数据集概述（": "数据集概述",
    "数据清洗（": "数据清洗",
    "数据转换（": "数据转换",
    "数据整合（": "数据整合",
    "模型选型（": "模型选型",
    "模型计算与验证（": "模型计算与验证",
    "可视化模板复现（": "可视化模板复现",
    "业务模型图（": "业务模型图",
    "交互式可视化大屏（": "交互式可视化大屏",
    "商业建议（": "商业建议",
    "AI 使用披露（": "AI 使用披露",
    "C 图表索引（": "C 图表索引",
}

SECTIONS = [
    ("引言", 1, [B("本章介绍研究背景、研究目的与研究意义，说明为什么选择 Olist 巴西电商数据，以及本项目要回答的核心商业问题。")]),
    ("研究背景", 2, [B(
        "近年来，巴西电商市场保持较快增长，Olist 作为连接中小卖家与消费者的平台，积累了订单、支付、物流、评价等多维数据。"
        "电商企业并不缺数据，缺的是把数据转化为可执行策略的能力：哪些客户贡献主要 GMV、哪些品类值得投入、复购为什么低、"
        "未来几个月销量如何。与此同时，课程要求将《Excel数据可视化——从图表到数据大屏》中的图表模板与 Python 商业分析模型结合，"
        "形成可复现的端到端项目，这为本文提供了明确的实践场景。"
    )]),
    ("研究目的", 2, [B(
        "本文旨在构建一条从公开数据获取、清洗治理、商业建模到可视化大屏的完整链路，回答四个问题：谁在买（RFM）、"
        "会不会复购（同期群）、卖什么（ABC 与波士顿矩阵）、未来卖多少（SARIMA），并用 Python 复现教材图表、"
        "搭建交互式大屏，使分析结论能够被业务人员直接使用。"
    )]),
    ("研究意义", 2, [B(
        "实践上，为中小电商提供低成本、可复现的客户与品类分析框架；方法上，将描述性分析（RFM、ABC、波士顿矩阵）"
        "与预测性分析（SARIMA）结合，避免只做静态报表；教学上，展示从 Excel 模板到 Python 可视化大屏的迁移路径，"
        "体现“数据获取—清洗—建模—可视化—洞察”的完整商业分析流程。"
    )]),

    ("数据准备与预处理", 1, [B("本章说明数据来源、合规获取方式、数据集结构以及清洗治理过程，为后续建模提供可复现的数据基础。")]),
    ("数据源介绍", 2, [
        B("数据来自 Kaggle 公开数据集 Brazilian E-Commerce Public Dataset by Olist，覆盖 2016-09-04 至 2018-10-17 的 "
          "99,441 个订单、112,650 条商品明细、99,224 条评价，共 9 张表。数据包含订单状态、客户、商品、卖家、支付、"
          "评价、物流时间与地理坐标，能够支撑多表关联、客户分层、品类分析和时间序列预测。"),
        T("表 1 数据表结构",
          ["表名", "行数", "粒度", "主键"],
          [
              ["orders", "99,441", "一行一个订单", "order_id"],
              ["order_items", "112,650", "一行一个订单商品", "order_id + order_item_id"],
              ["payments", "103,886", "一行一笔支付", "order_id + payment_sequential"],
              ["reviews", "99,224", "一行一条评价", "review_id（需校验）"],
              ["customers", "99,441", "一行一个订单客户", "customer_id"],
              ["products", "32,951", "一行一个商品", "product_id"],
              ["sellers", "3,095", "一行一个卖家", "seller_id"],
              ["geolocation", "1,000,163", "一行一个邮编坐标", "无唯一键"],
              ["category_translation", "71", "一行一个品类翻译", "product_category_name"],
          ], widths=[4.2, 2.8, 3.2, 4.4], font_size=9),
    ]),
    ("合规获取方式", 2, [B(
        "本项目使用 Kaggle 官方数据集 API 获取数据，属于课程允许的“官方下载文件 / 官方数据接口”方式。该接口对公开数据集返回 302 "
        "跳转到 Google Storage 的签名地址，脚本使用 Python 标准库 urllib 跟随跳转，没有编写或使用任何网页爬虫。下载脚本内置 "
        "SHA256 校验（967e41e04fc306fe604e2a693f488995a8b41e5047418f8a5c8e4abd6deca784），保证文件完整、可复现。"
    )]),
    ("数据来源合规声明", 2, [B(
        "数据集采用 CC BY-NC-SA 4.0 许可，仅用于课程作业与学术研究，需署名、非商业使用并以相同方式共享。原始数据体积较大且受许可约束，"
        "因此不提交到公开仓库；仓库只保留下载脚本、校验信息与数据处理代码，任何人克隆仓库后运行下载脚本即可复现。"
    )]),
    ("数据集概述", 2, [B(
        "orders 为主表，通过 customer_id 关联 customers，通过 order_id 关联 order_items、payments、reviews；order_items 通过 "
        "product_id、seller_id 关联 products、sellers；products.product_category_name 关联品类翻译表。主外键检查显示 "
        "orders→customers、items→orders/products/sellers、payments/reviews→orders 的命中率均为 100%。99,441 个订单中有 "
        "98,666 个含商品明细，775 个无明细，主要为取消或不可用订单。"
    )]),

    ("数据清洗与治理", 1, [B("本章说明缺失值、异常值、重复值、类型统一、指标派生与多表整合的处理规则。")]),
    ("数据清洗", 2, [B(
        "时间字段统一转换为 datetime；邮编前缀按字符串读取并补足 5 位，避免 01037 被解析为 1037；商品品类缺失标记为 unknown，"
        "重量与长宽高缺失用中位数填充并增加 has_dimensions 标记；评价按订单取最新一条，review_id 重复的记录保留并加标记；"
        "地理表按邮编前缀聚合，经纬度取中位数，1,000,163 行压缩为 19,015 行。"
    )]),
    ("数据转换", 2, [B(
        "派生 purchase_month、delivery_days、late_days、is_late、approval_hours、product_volume_cm3、item_count、gmv 等指标。"
        "is_late 仅在“实际送达时间”和“承诺送达时间”均存在时计算，避免把缺失值误判为准时；金额类指标统一为 BRL。"
    )]),
    ("数据整合", 2, [B(
        "生成订单级宽表 orders_enriched（99,441 行 × 36 列），整合商品、支付、评价、客户信息；导出 CSV 与带索引的 SQLite 数据库。"
        "本项目使用全量数据，不做抽样；所有随机过程固定 RANDOM_SEED = 42，保证结果可复现。"
    )]),

    ("商业分析模型", 1, [B("本章介绍模型选型、参数与阈值依据，以及各模型的计算结果与验证指标。")]),
    ("模型选型", 2, [B(
        "选取 5 个核心模型和 1 个辅助分析：RFM 客户价值分层、同期群留存、ABC/Pareto 商品分类、波士顿矩阵品类投入、"
        "SARIMA 月度 GMV 预测，以及订单转化漏斗。RFM 回答“谁贡献了 GMV”，同期群回答“客户会不会复购”，ABC 与波士顿矩阵回答"
        "“哪些商品/品类值得投入”，SARIMA 回答“未来卖多少”，漏斗回答“履约链路是否稳定”。"
    )]),
    ("参数与阈值设定依据", 2, [B(
        "RFM 以 customer_unique_id 为客户唯一标识，快照日为 2018-09-04；R 与 M 用中位数切分（R 中位数 225 天，M 中位数 R$89.89），"
        "F 中位数为 1，因此以 F≥2 作为复购阈值。选用中位数而非均值，是因为 GMV 右偏严重，均值会被少数大客户拉高。"
        "ABC 以累计 GMV 80%/95% 作为 A/B/C 边界。波士顿矩阵同比统一取 2017 年 1—8 月 vs 2018 年 1—8 月，"
        "避免 2018 年数据不完整；以份额中位数 0.49%、增长率中位数 188% 划分四象限。SARIMA 使用 2017-01—2018-08 共 20 个月，"
        "前 14 个月训练、后 6 个月回测，按回测 RMSE 选择模型。"
    )]),
    ("模型计算与验证", 2, [
        B("RFM 将 94,983 位客户分为 8 类。重要发展客户 22,704 人（23.90%）贡献 39.08% GMV，重要挽留客户 22,385 人（23.57%）"
          "贡献 38.93%，两者合计 47.5% 的客户贡献 78.0% 的 GMV；重要价值/保持客户仅 2,403 人，但平均订单数约 2.1，是复购核心。"
          "表 2 给出八分层结果，图 17 为对应的 GMV 贡献图。"),
        T("表 2 RFM 八分层结果",
          ["分层", "客户数", "客户占比", "GMV 占比"],
          [
              ["重要发展客户", "22,704", "23.90%", "39.08%"],
              ["重要挽留客户", "22,385", "23.57%", "38.93%"],
              ["一般挽留客户", "23,683", "24.93%", "8.29%"],
              ["一般发展客户", "23,324", "24.56%", "8.14%"],
              ["重要价值客户", "1,340", "1.41%", "2.98%"],
              ["重要保持客户", "1,063", "1.12%", "2.36%"],
              ["一般价值客户", "243", "0.26%", "0.11%"],
              ["一般保持客户", "241", "0.25%", "0.11%"],
          ], widths=[3.8, 3.2, 3.2, 3.2]),
        F(FIG_B / "fig17_RFM分层贡献.png", "图 17 RFM 八分层 GMV 贡献"),
        B("同期群分析共识别 23 个首购月同期群，其中 21 个样本量≥30。第 1 月平均留存 0.44%，第 3 月 0.21%，第 6 月 0.19%，"
          "说明 Olist 的复购率极低，增长主要依赖新客。图 18 为同期群留存热力图。"),
        F(FIG_B / "fig18_同期群留存热力图.png", "图 18 同期群留存热力图"),
        B("ABC 分类中，A 类 8,478 个商品（25.9%）贡献 80.00% GMV；B 类 11,226 个（34.3%）贡献 15.00%；C 类 13,025 个（39.8%）"
          "贡献 5.00%。前 20 个单品仅贡献 5.4%，说明 A 类是一个“长尾中的头部集合”，不能只盯少数爆款。表 3 与图 19 给出分类结果。"),
        T("表 3 ABC 分类结果",
          ["类别", "商品数", "商品占比", "GMV 占比"],
          [["A", "8,478", "25.90%", "80.00%"], ["B", "11,226", "34.30%", "15.00%"], ["C", "13,025", "39.80%", "5.00%"]],
          widths=[2.5, 3.5, 3.5, 3.5]),
        F(FIG_B / "fig19_ABC帕累托图.png", "图 19 ABC/Pareto 图"),
        B("波士顿矩阵共有 63 个品类参与划分：明星 15 个（health_beauty、watches_gifts、housewares、baby、telephony 等），"
          "现金牛 17 个（bed_bath_table、sports_leisure、computers_accessories 等），问题 17 个，瘦狗 14 个。"
          "表 4 给出象限分布，图 20 为四象限气泡图。"),
        T("表 4 波士顿矩阵象限分布",
          ["象限", "品类数", "代表品类"],
          [
              ["明星（高份额/高增长）", "15", "health_beauty、watches_gifts、housewares、baby"],
              ["现金牛（高份额/低增长）", "17", "bed_bath_table、sports_leisure、computers_accessories"],
              ["问题（低份额/高增长）", "17", "需结合利润率与复购判断"],
              ["瘦狗（低份额/低增长）", "14", "考虑收缩 SKU"],
          ], widths=[4.6, 2.6, 7.4]),
        F(FIG_B / "fig20_波士顿矩阵.png", "图 20 波士顿矩阵"),
        B("SARIMA 部分比较了 ARIMA(1,1,1)、ARIMA(0,1,1) 与多个 SARIMA 候选。季节差分模型因月度样本仅 20 个月而失败，"
          "最终按回测 RMSE 选择 ARIMA(1,1,1)：RMSE 67,582，MAPE 5.59%，显著优于季节朴素基准（MAPE 50.56%）。"
          "未来 6 个月预测均值约 R$861,872/月，并给出 95% 置信区间。表 5 与图 16 为模型比较与预测结果。"),
        T("表 5 SARIMA 候选模型回测结果",
          ["模型", "RMSE (R$)", "MAPE"],
          [["ARIMA(1,1,1)", "67,582", "5.59%"], ["ARIMA(0,1,1)", "94,807", "7.43%"],
           ["SARIMA(1,1,1)(1,0,1,12)", "75,152", "6.04%"], ["SARIMA(0,1,1)(0,0,1,12)", "114,698", "9.88%"],
           ["季节朴素基准", "489,873", "50.56%"]], widths=[6.5, 4.0, 3.0]),
        F(FIG_B / "fig16_月度GMV与预测.png", "图 16 月度 GMV 与六个月预测"),
        B("订单转化漏斗为：下单 99,441 → 支付通过 99,281 → 发货 97,658 → 送达 96,476，总体送达率 97.0%。"
          "主要损失来自取消/不可用订单，支付与发货环节损失较小。表 6 与图 25 给出漏斗结果。"),
        T("表 6 订单转化漏斗",
          ["阶段", "订单数", "占下单", "较上一阶段"],
          [["下单", "99,441", "100.0%", "—"], ["支付通过", "99,281", "99.8%", "99.8%"],
           ["发货", "97,658", "98.2%", "98.4%"], ["送达", "96,476", "97.0%", "98.8%"]],
          widths=[3.5, 3.5, 3.5, 3.5]),
        F(FIG_B / "fig25_订单转化漏斗.png", "图 25 订单转化漏斗"),
    ]),

    ("Python 数据分析与可视化", 1, [B("本章展示探索性分析、教材模板复现、业务模型图与交互式大屏的实现与结果。")]),
    ("数据探索性分析", 2, [B(
        "数据体检显示：订单表缺失率 0.62%，评价表缺失率 21.01%（主要是评论正文），商品表缺失率 0.83%，地理表无缺失但有 261,831 条全行重复。"
        "GMV 合计 R$13,591,643.70，运费合计 R$2,251,909.54；支付金额与“商品金额+运费”的平均绝对差仅 R$0.03，说明支付与订单明细一致性良好。"
        "有效订单（非取消/不可用且有商品明细）共 98,199 个。"
    )]),
    ("可视化模板复现", 2, [
        B("复现教材第二章前 15 套模板：渐变柱形、带均值柱形、渐变圆角柱形、标注柱形、层叠柱形、蝴蝶图、REPT 蝴蝶、数值百分比、"
          "对比柱形、甘特图、平滑折线、菱形走势、对比折线、单值圆环、水球图。方法是从 xlsx 的 chart XML 解析颜色、字号、渐变、"
          "布局与标注位置，用 matplotlib 逐像素还原；15 张图平均通道误差 1.61—5.76/255，像素一致率 87.42%—96.43%，"
          "交付图为 200 dpi。图 1、图 5、图 8、图 10、图 14 为代表性示例，其余见附录 C。"),
        F(FIG_T / "fig01_渐变柱形图.png", "图 1 渐变柱形图（模板复现）"),
        F(FIG_T / "fig05_层叠柱形图.png", "图 5 层叠柱形图（模板复现）"),
        F(FIG_T / "fig08_数值百分比.png", "图 8 数值百分比（模板复现）"),
        F(FIG_T / "fig10_甘特图.png", "图 10 甘特图（模板复现）"),
        F(FIG_T / "fig14_单值圆环图.png", "图 14 单值圆环图（模板复现）"),
    ]),
    ("业务模型图", 2, [
        B("10 张业务图与模型一一对应：图 16 月度 GMV 与预测、图 17 RFM、图 18 同期群、图 19 ABC、图 20 波士顿、图 21 品类 Top 15、"
          "图 22 配送延迟与评分、图 23 支付方式、图 24 RFM 气泡散点、图 25 订单漏斗。前文已展示图 16—20 与图 25，这里补充图 21—24。"),
        F(FIG_B / "fig21_品类GMV前15.png", "图 21 2018 年 1—8 月 GMV 前 15 品类"),
        F(FIG_B / "fig22_配送延迟与评价.png", "图 22 配送延迟与评价"),
        F(FIG_B / "fig23_支付方式结构.png", "图 23 支付方式结构"),
        F(FIG_B / "fig24_RFM气泡散点图.png", "图 24 RFM 气泡散点"),
    ]),
    ("交互式可视化大屏", 2, [
        B("使用 Streamlit + Plotly 搭建交互式大屏，包含 5 个 KPI（GMV、订单数、客户数、客单价、延迟送达率）、4 个页签"
          "（经营总览、客户与留存、品类与预测、数据明细）、日期/州/品类/订单状态筛选、明细表与 CSV 下载。"
          "RFM、同期群、波士顿矩阵、SARIMA 为离线全量结果，其余图表随筛选联动。运行命令为 "
          "streamlit run scripts/06_dashboard.py，默认地址 http://localhost:8501。图 26 为大屏截图。"),
        F(DASH, "图 26 Streamlit 交互式大屏", width=14.5),
    ]),

    ("商业案例分析", 1, [B("本章结合模型结果给出商业解读与可执行建议。")]),
    ("案例背景", 2, [B(
        "Olist 连接巴西中小卖家与消费者，2016—2018 年处于快速增长期。平台的核心问题是客户复购率低、品类分散、物流时效差异大，"
        "需要通过数据分析识别高价值客户、核心品类与履约短板。"
    )]),
    ("结果解读", 2, [B(
        "第一，客户高度集中：47.5% 的高价值客户贡献 78.0% GMV，其中“重要挽留”客户平均 R 约 372 天，存在明显流失风险。"
        "第二，复购率低：第 1 月留存仅 0.44%，第 6 月 0.19%，增长主要依赖新客。第三，商品贡献高度集中：25.9% 的商品贡献 80% GMV，"
        "但前 20 个单品仅贡献 5.4%，说明 A 类商品数量多、长尾明显。第四，品类分化：15 个明星品类、17 个现金牛、17 个问题、14 个瘦狗。"
        "第五，物流是体验分水岭：准时/提前订单平均评分 4.29，4—7 天降至 2.32，15 天以上仅 1.71，差评率 78.8%。"
        "第六，支付以信用卡（78.5%）和 Boleto（18.0%）为主。第七，未来 6 个月 GMV 预测均值约 R$86 万/月，MAPE 5.59%。"
    )]),
    ("商业建议", 2, [
        B("（1）客户分层运营：对“重要挽留”客户做 30/60/90 天召回，对“重要发展”客户发二次购买券，对“重要价值/保持”客户提供会员等级与专属客服。"),
        B("（2）提升复购：首购后 30 天触发品类交叉推荐，60 天触发免运费券，90 天触发老客专属折扣；按品类复购周期设计触达节奏。"),
        B("（3）商品与库存：A 类商品保库存、保曝光；B 类做组合销售；C 类做长尾清理。不要只盯少数爆款，A 类是 8,478 个 SKU 的集合。"),
        B("（4）品类策略：明星品类加大投放；现金牛品类维持份额、控制成本；问题品类小步测试；瘦狗品类收缩 SKU。"),
        B("（5）物流体验：按承运商与州考核时效；对预计延迟订单主动通知与补偿；把承诺送达时间纳入商品页，降低“延迟 4 天以上”的差评。"),
        B("（6）支付结构：保留 Boleto，优化信用卡分期展示；用代金券做召回与清库存，提升借记卡与代金券的使用率。"),
        B("（7）预测与备货：用预测中值做预算、用 95% 区间上限做备货与人力弹性，每月滚动更新，避免只按点预测备货。"),
        B("（8）履约漏斗：对取消原因分类（支付失败、库存不可用、客户主动取消），优先修复可运营环节，把履约稳定性纳入承运商与库存考核。"),
    ]),

    ("结论与展望", 1, [B("本章总结研究结论、研究局限与后续改进方向，并披露 AI 工具的使用情况。")]),
    ("研究结论", 2, [B(
        "本文完成了 9 表关联、99,441 个订单的清洗与治理，构建了订单级宽表；使用 RFM、同期群、ABC/Pareto、波士顿矩阵、SARIMA "
        "和订单漏斗六个模型，得到客户集中、复购低、A 类商品贡献 80%、物流延迟显著影响评分、预测 MAPE 5.59% 等结论；"
        "复现教材 15 张模板图、绘制 10 张业务图并搭建 Streamlit 大屏，形成了可复现的端到端商业分析项目。"
    )]),
    ("局限性与展望", 2, [B(
        "研究仍有局限：数据仅覆盖 2016—2018 年，月度样本只有 20 个月，SARIMA 的季节项容易过拟合；数据缺少成本与利润字段，"
        "无法计算 CLV 与真实利润；RFM 阈值依赖中位数，换数据后需要重新校准。未来可引入利润与成本数据做客户终身价值分析，"
        "用 Prophet、LSTM 与 SARIMA 对比，接入实时数据做在线大屏，并通过 A/B 实验验证运营策略。"
    )]),
    ("AI 使用披露", 2, [B(
        "本文使用 Claude Code（Anthropic）辅助完成数据下载脚本、清洗流水线、RFM/同期群/ABC/波士顿/SARIMA 代码、matplotlib/Plotly "
        "图表、Streamlit 大屏与论文框架。关键提示词包括“基于 Olist 数据集设计分析流程并说明阈值依据”“用 matplotlib 绘制 10 张业务图，"
        "配色使用通过 CVD 校验的调色板”“生成包含 6 个以上联动组件的 Streamlit 大屏”。作者检查了数据许可、字段口径、模型阈值、"
        "图表结论与论文表述；所有数字均来自可复现脚本，未使用 AI 编造数据。AI 仅作为辅助工具，论文观点与最终表述由作者负责。"
    )]),
]


REFERENCES = [
    "Olist. Brazilian E-Commerce Public Dataset by Olist[DB/OL]. Kaggle, 2018. "
    "https://www.kaggle.com/datasets/olistbr/brazilian-ecommerce.",
    "郭宏远. Excel数据可视化——从图表到数据大屏[M]. 北京: 清华大学出版社, 2022.",
    "Hughes A M. Strategic Database Marketing[M]. Chicago: Probus Publishing, 1994.",
    "Chen D, Sain S L, Guo K. Data mining for the online retail industry: A case study of "
    "RFM model-based customer segmentation[J]. Journal of Database Marketing & Customer Strategy "
    "Management, 2012, 19(3): 197-208.",
    "Box G E P, Jenkins G M, Reinsel G C, et al. Time Series Analysis: Forecasting and Control[M]. "
    "5th ed. Hoboken: Wiley, 2015.",
    "Streamlit. Streamlit Documentation[EB/OL]. https://docs.streamlit.io.",
    "Plotly. Plotly Python Open Source Graphing Library[EB/OL]. https://plotly.com/python/.",
]

CHART_INDEX = [
    ["图 1—15", "outputs/figures/templates/fig01—fig15", "教材模板复现（15 张，200 dpi）"],
    ["图 16", "fig16_月度GMV与预测.png", "月度 GMV 与六个月预测"],
    ["图 17", "fig17_RFM分层贡献.png", "RFM 八分层 GMV 贡献"],
    ["图 18", "fig18_同期群留存热力图.png", "同期群留存热力图"],
    ["图 19", "fig19_ABC帕累托图.png", "ABC/Pareto 图"],
    ["图 20", "fig20_波士顿矩阵.png", "波士顿矩阵"],
    ["图 21", "fig21_品类GMV前15.png", "2018 年 1—8 月 GMV 前 15 品类"],
    ["图 22", "fig22_配送延迟与评价.png", "配送延迟与评价"],
    ["图 23", "fig23_支付方式结构.png", "支付方式结构"],
    ["图 24", "fig24_RFM气泡散点图.png", "RFM 气泡散点"],
    ["图 25", "fig25_订单转化漏斗.png", "订单转化漏斗"],
    ["图 26", "dashboard_preview.png", "Streamlit 交互式大屏"],
]


def set_run(run, name: str, size: float, bold: bool = False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def set_para_font(p, name: str, size: float, bold: bool = False):
    for run in p.runs:
        set_run(run, name, size, bold)


def clear_para(p):
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def new_after(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def add_body(cursor: Paragraph, text: str) -> Paragraph:
    p = new_after(cursor)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    set_run(run, "宋体", 12)
    return p


def add_caption(cursor: Paragraph, text: str) -> Paragraph:
    p = new_after(cursor)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run(run, "宋体", 10.5)
    return p


def add_table(doc, cursor: Paragraph, caption: str, headers: list[str],
              rows: list[list[str]], widths: list[float] | None = None,
              font_size: float = 10.5) -> Paragraph:
    cap = add_caption(cursor, caption)
    cap.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if table._tbl.tblPr.find(qn("w:tblLayout")) is None:
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        table._tbl.tblPr.append(layout)
    cap._p.addnext(table._tbl)
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(text)
        set_run(run, "宋体", font_size, bold=True)
        hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 表头跨页重复
    trPr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trPr.append(tbl_header)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(text)
            set_run(run, "宋体", font_size)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        trPr.append(cant_split)
    if widths:
        for i, width in enumerate(widths):
            table.columns[i].width = Cm(width)
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    after = OxmlElement("w:p")
    table._tbl.addnext(after)
    return Paragraph(after, cursor._parent)


def add_figure(cursor: Paragraph, path: Path, caption: str, width_cm: float) -> Paragraph:
    p = new_after(cursor)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    if path.is_file():
        run = p.add_run()
        run.add_picture(str(path), width=Cm(width_cm))
    else:
        run = p.add_run(f"【图片缺失：{path}】")
        set_run(run, "宋体", 10.5)
    return add_caption(p, caption)


def find_heading(doc, prefix: str, level: int | None):
    for p in doc.paragraphs:
        if level is not None and p.style.name != f"Heading {level}":
            continue
        if p.text.strip().startswith(prefix):
            return p
    return None


def fill_cover(doc):
    cover = {
        3: f"题　　目：{TITLE}",
        4: "学　　院：数据科学学院",
        5: "专业班级：________________",
        6: "学　　号：________________",
        7: f"姓　　名：{AUTHOR}",
        8: f"完成日期：{DATE}",
    }
    for idx, text in cover.items():
        p = doc.paragraphs[idx]
        clear_para(p)
        run = p.add_run(text)
        set_run(run, "楷体_GB2312", 14)

    p = doc.paragraphs[11]
    clear_para(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(TITLE)
    set_run(run, "黑体", 16, bold=True)

    p = doc.paragraphs[12]
    clear_para(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(AUTHOR), "宋体", 10.5)

    p = doc.paragraphs[13]
    clear_para(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(COLLEGE), "宋体", 10.5)

    # 先保存引用，插入摘要正文后段落索引会变化
    p_abs = doc.paragraphs[15]
    p_kw = doc.paragraphs[16]
    p_note = doc.paragraphs[17]

    # 摘要标题 + 正文
    clear_para(p_abs)
    set_run(p_abs.add_run("摘　要"), "黑体", 14)
    add_body(p_abs, ABSTRACT)

    # 关键词
    clear_para(p_kw)
    set_run(p_kw.add_run("关键词："), "黑体", 14)
    set_run(p_kw.add_run(KEYWORDS), "宋体", 12)

    # 删除红色提示
    p_note._element.getparent().remove(p_note._element)


def write_markdown():
    lines = [f"# {TITLE}", "", f"作者：{AUTHOR}　　{COLLEGE}", "", "## 摘要", "", ABSTRACT, "", f"**关键词：** {KEYWORDS}", ""]
    for prefix, level, blocks in SECTIONS:
        lines.append(f"{'#' * (level + 1)} {prefix}")
        lines.append("")
        for block in blocks:
            if block[0] == "p":
                lines.append(block[1])
                lines.append("")
            elif block[0] == "table":
                _, cap, headers, rows, _widths, _font = block
                lines.append(f"**{cap}**")
                lines.append("")
                lines.append("| " + " | ".join(headers) + " |")
                lines.append("|" + "|".join(["---"] * len(headers)) + "|")
                for row in rows:
                    lines.append("| " + " | ".join(row) + " |")
                lines.append("")
            elif block[0] == "figure":
                _, path, cap, _ = block
                try:
                    rel = path.relative_to(ROOT).as_posix()
                except ValueError:
                    rel = str(path)
                lines.append(f"![{cap}]({rel})")
                lines.append("")
    lines.append("## 参考文献")
    lines.append("")
    for i, ref in enumerate(REFERENCES, 1):
        lines.append(f"[{i}] {ref}")
        lines.append("")
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")


def clean_headings(doc):
    """去掉模板标题里的评分提示括号，保留干净标题。"""
    for p in doc.paragraphs:
        text = p.text.strip()
        for prefix, clean in HEADING_CLEAN.items():
            if text.startswith(prefix.strip()):
                clear_para(p)
                run = p.add_run(clean)
                if p.style.name == "Heading 1":
                    set_run(run, "黑体", 16, bold=True)
                elif p.style.name == "Heading 2":
                    set_run(run, "黑体", 14, bold=True)
                else:
                    set_run(run, "宋体", 12, bold=True)
                p.paragraph_format.keep_with_next = True
                break


def remove_empty_paragraphs(doc):
    """删除正文中的空段落，压缩模板自带的空白；保留封面/摘要和分页符。"""
    paragraphs = doc.paragraphs
    start = 0
    for i, p in enumerate(paragraphs):
        if p.style.name == "Heading 1" and p.text.strip().startswith("引言"):
            start = i
            break
    for p in list(paragraphs)[start:]:
        if p.text.strip():
            continue
        if p._p.find(".//" + qn("w:drawing")) is not None:
            continue
        has_page_break = any(
            br.get(qn("w:type")) == "page" for br in p._p.iter(qn("w:br"))
        )
        if has_page_break:
            continue
        if p._p.find(".//" + qn("w:sectPr")) is not None:
            continue
        p._element.getparent().remove(p._element)


def main() -> int:
    doc = docx.Document(str(TEMPLATE))
    fill_cover(doc)
    clean_headings(doc)
    # 删除模板自带的示例参考文献
    for p in list(doc.paragraphs):
        if p.text.strip().startswith("[1] ***"):
            p._element.getparent().remove(p._element)
            break

    # 删除模板末尾的评分表
    if doc.tables:
        table = doc.tables[0]
        table._element.getparent().remove(table._element)

    # 按标题插入内容
    for prefix, level, blocks in SECTIONS:
        heading = find_heading(doc, prefix, level)
        if heading is None:
            print(f"警告: 未找到标题 {prefix!r}", file=sys.stderr)
            continue
        cursor = heading
        for block in blocks:
            if block[0] == "p":
                cursor = add_body(cursor, block[1])
            elif block[0] == "table":
                _, cap, headers, rows, widths, font_size = block
                cursor = add_table(doc, cursor, cap, headers, rows, widths, font_size)
            elif block[0] == "figure":
                _, path, cap, width = block
                cursor = add_figure(cursor, path, cap, width)

    # 参考文献
    ref_heading = find_heading(doc, "参考文献", 1)
    if ref_heading is not None:
        cursor = ref_heading
        for i, ref in enumerate(REFERENCES, 1):
            p = new_after(cursor)
            p.paragraph_format.line_spacing = 1.5
            set_run(p.add_run(f"[{i}] {ref}"), "宋体", 10.5)
            cursor = p

    # 附录
    appendix_map = {
        "A 核心代码片段": [
            "import pandas as pd",
            "from statsmodels.tsa.statespace.sarimax import SARIMAX",
            "",
            "snapshot = valid['purchase_date'].max() + pd.Timedelta(days=1)",
            "rfm = valid.groupby('customer_unique_id').agg(",
            "    last_purchase=('purchase_date', 'max'),",
            "    frequency=('order_id', 'nunique'),",
            "    monetary=('gmv', 'sum'),",
            ")",
            "rfm['recency'] = (snapshot - rfm['last_purchase']).dt.days",
            "",
            "model = SARIMAX(train, order=(1, 1, 1), seasonal_order=(0, 0, 0, 0),",
            "                enforce_stationarity=False, enforce_invertibility=False)",
            "fit = model.fit(disp=False)",
            "forecast = fit.get_forecast(steps=6)",
        ],
        "B Gitee/github仓库链接与目录结构": [
            "仓库链接：https://github.com/<your-account>/olist-analysis（提交前替换为真实公开链接）",
            "",
            "olist-analysis/",
            "├─ data/raw/           # 原始数据（git 忽略）",
            "├─ data/clean/         # 清洗数据与 SQLite（git 忽略）",
            "├─ scripts/            # 00—08 分析脚本",
            "├─ outputs/figures/    # 模板图、业务图、大屏截图",
            "├─ outputs/tables/     # 模型结果 CSV",
            "├─ docs/               # 数据说明、体检/清洗/模型报告、论文",
            "├─ requirements.txt",
            "├─ run_all.py",
            "└─ README.md",
        ],
    }
    for prefix, lines in appendix_map.items():
        p = find_heading(doc, prefix, None)
        if p is None:
            print(f"警告: 未找到附录标题 {prefix!r}", file=sys.stderr)
            continue
        cursor = p
        for line in lines:
            np = new_after(cursor)
            np.paragraph_format.line_spacing = 1.0
            run = np.add_run(line)
            set_run(run, "Consolas" if line.startswith(("import", "from", "snapshot", "rfm", "model", "fit", "forecast", ")", "    ")) else "宋体", 9)
            cursor = np

    # 图表索引
    p = find_heading(doc, "C 图表索引", None)
    if p is not None:
        add_table(doc, p, "表 C1 图表索引", ["图号", "文件", "说明"], CHART_INDEX,
                  widths=[2.8, 7.6, 4.2])

    remove_empty_paragraphs(doc)
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(str(OUT_DOCX))
    except PermissionError:
        print(
            f"错误: 无法写入 {OUT_DOCX}\n"
            "该文件可能正在 Word 中打开，请关闭后重试。",
            file=sys.stderr,
        )
        return 1
    write_markdown()
    print(f"已生成 Word 初稿: {OUT_DOCX}")
    print(f"已生成 Markdown: {OUT_MD}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
